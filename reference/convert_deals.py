# -*- coding: utf-8 -*-
"""
convert_deals.py (v2) – Chuyển file export Deals từ CRM (raw) sang file template phân tích.

Nguồn dữ liệu: sheet "1 - Bảng danh mục sản phẩm dự án" của file export CRM.
  - Dòng có ID  = dòng deal (đủ thông tin deal + dòng BOM đầu tiên).
  - Dòng ID trống = dòng BOM nối tiếp của deal phía trên.
  - BOM chỉ lấy từ 9 cột "Bảng danh mục sản phẩm dự án - …". Cột text tổng hợp
    "Bảng danh mục sản phẩm dự án" KHÔNG được dùng.

Cách chạy:
    python convert_deals.py <file_raw.xlsx> <file_output.xlsx> [--ref-date dd/mm/yyyy]

--ref-date: ngày chốt số liệu (tính quá hạn forecast, close date đã qua).
            Mặc định lấy từ tên file (..._dd_mm_yy.xlsx), nếu không có thì lấy ngày hôm nay.

Yêu cầu: pandas, openpyxl.
"""
import sys, re, html, unicodedata, argparse
from datetime import datetime
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

# ============================== THAM SỐ ==============================
SHEET_PREFIX = '1 - Bảng danh mục'   # tên sheet bị Excel cắt còn 31 ký tự nên so theo tiền tố
USD_MISLABEL_THRESHOLD = 1_000_000   # deal gắn USD có giá trị >= ngưỡng này coi là VND nhập nhầm nhãn
USD_RATE = 26_000                    # GIẢ ĐỊNH, chỉ dùng cho deal USD < ngưỡng (hiện chưa có deal nào)
BOM_TOLERANCE = 0.05                 # lệch giá trị so với tổng BOM > 5% thì gắn cờ
DEAL_SIZE = [(1e9, '< 1B'), (2e9, '1B - 2B'), (5e9, '2B - 5B')]   # >= 5B -> '> 5B'
TEST_NAMES = re.compile(r'(^|[^a-z])test([^a-z]|$)|\bko sale\b|\bkhong sale\b')
LEAD_STAGE = 'Đăng kí cơ hội'

B = 'Bảng danh mục sản phẩm dự án - '
BOM_COLS = {k: B + k for k in ['Hãng', 'Tên giải pháp', 'Số lượng', 'Đơn giá', 'Thành tiền',
                               'Thời hạn bản quyền', 'Ghi chú', 'Sales user', 'Link deal']}
BOM_CORE = [BOM_COLS[k] for k in ['Hãng', 'Tên giải pháp', 'Số lượng', 'Đơn giá', 'Thành tiền']]
ACTIVITY = ['Số email', 'Số cuộc họp', 'Số cuộc gọi', 'Số ghi chú', 'Số ghi nhanh']

# Cột bắt buộc phải có trong sheet nguồn (kiểm tra trước khi chạy)
REQUIRED = ['ID', 'Tên thương vụ', 'Giá trị thương vụ', 'Tiền tệ thương vụ', 'Chủ sở hữu thương vụ',
            'Giai đoạn', 'Trạng thái thương vụ', 'Tên công ty', 'Phân loại HĐ', 'Reseller', 'Tên hãng',
            'Thông tin BOM (model, số lượng, năm support...)', 'Tên dự án', 'Timeline dự án',
            'Ngày ký hợp đồng', 'Phân loại nhóm deal', 'Lý do thất bại', 'Failed additional comment',
            'Từ lúc', 'Cập nhật lần cuối'] + ACTIVITY + list(BOM_COLS.values())

VENDOR_MAP = {
    'KASPERSKY': 'Kaspersky', 'KAPERSKY': 'Kaspersky', 'GTB': 'GTB', 'SOPHOS': 'Sophos', 'OPSWAT': 'OPSWAT',
    'NETSCOUT': 'NetScout', 'DELINEA': 'Delinea', 'CLOUDFLARE': 'Cloudflare', 'PROGRESS': 'Progress',
    'INFOEXPRESS': 'Infoexpress', 'OPENTEXT': 'OpenText', 'BARRACUDA': 'Barracuda', 'RADWARE': 'Radware',
    'SAFETICA': 'Safetica', 'SAFEBREACH': 'SafeBreach', 'HCL SOFTWARE': 'HCL Software', 'HCL': 'HCL Software',
    'ZECURION': 'Zecurion', 'ACRONIS': 'Acronis', 'QUALYS': 'Qualys', 'STELLAR CYBER': 'Stellar Cyber',
    'PENTA SECURITY': 'Penta Security', 'ARCON': 'Arcon', 'PENTERA': 'Pentera', 'PENTARA': 'Pentera',
    'TRUSTWAVE': 'Trustwave', 'NETGEAR': 'Netgear', 'THALES': 'Thales', 'FORCEPOINT': 'Forcepoint',
    'SECPOD': 'SecPod', 'PROOFPOINT': 'Proofpoint', 'CYBLE': 'Cyble', 'FORTRA': 'Fortra', 'VIAVI': 'Viavi',
    'TRELLIX': 'Trellix', 'EFFICIENTIP': 'EfficientIP', 'HÃNG KHÁC': 'Khác', 'SẢN PHẨM KHÁC': 'Khác', 'KHÁC': 'Khác',
}
# Từ khóa để suy ra hãng khi BOM không có "Hãng:" (tìm trên text đã bỏ dấu, chữ thường)
VENDOR_KEYWORDS = [
    ('Kaspersky', r'kaspersky|kapersky|\bkas\b'), ('Sophos', r'sophos'), ('GTB', r'\bgtb\b'),
    ('OPSWAT', r'opswat|metadefender'), ('NetScout', r'netscout|arbor'), ('Delinea', r'delinea'),
    ('Cloudflare', r'cloudflare'), ('Progress', r'progress|whatsup|\bwug\b|loadmaster|kemp'),
    ('Infoexpress', r'infoexpress|easynac'), ('OpenText', r'opentext|arcsight|fortify|\bsmax\b|\bopds\b|voltage|encase'),
    ('Barracuda', r'barracuda'), ('Radware', r'radware'), ('Safetica', r'safetica'), ('SafeBreach', r'safebreach'),
    ('HCL Software', r'\bhcl\b|appscan|\bhlc\b'), ('Zecurion', r'zecurion'), ('Acronis', r'acronis'),
    ('Qualys', r'qualys'), ('Stellar Cyber', r'stellar'), ('Penta Security', r'penta'), ('Arcon', r'arcon'),
    ('Pentera', r'pentera|pentara'), ('Trellix', r'trellix'), ('SecPod', r'secpod'), ('Thales', r'thales'),
    ('Proofpoint', r'proofpoint'), ('Forcepoint', r'forcepoint'), ('Trustwave', r'trustwave'),
    ('Netgear', r'netgear'), ('Viavi', r'viavi'), ('EfficientIP', r'efficient ?ip'),
    ('Cyble', r'cyble'), ('Fortra', r'fortra'),
]
# Từ khóa nhận diện phần "hãng/giải pháp" trong tên thương vụ
SOLUTION_KW = {'waf', 'ddos', 'dr', 'fw', 'firewall', 'sophos', 'kaspersky', 'kapersky', 'kas', 'gtb', 'safetica',
               'safebreach', 'smax', 'loadmaster', 'opswat', 'qualys', 'progress', 'pam', 'secpod', 'renew',
               'foundations', 'optimum', 'appscan', 'hlc', 'endpoint', 'dlp', 'nac', 'easynac', 'trellix', 'delinea',
               'netscout', 'opds', 'cloudflare', 'opentext', 'barracuda', 'radware', 'infoexpress', 'acronis',
               'zecurion', 'stellar', 'pentera', 'fortify', 'arcsight', 'wug', 'whatsup', 'hcl', 'kemp', 'edr',
               'siem', 'mfa', 'ztna', 'appsec', 'db', 'security'}
FILLER = {'va', 'and', 'plus', 'cap', 'do', 'bo', 'sung'}
PROJECT_PREFIX = re.compile(r'^(mua sam|giai phap|du an|trien khai|xay dung|nang cap|thay the)')



# ============================== HÀM TIỆN ÍCH ==============================
def fold(s):
    """chữ thường, bỏ dấu tiếng Việt – dùng để so khớp"""
    s = unicodedata.normalize('NFD', str(s).lower()).replace('đ', 'd')
    return ''.join(c for c in s if unicodedata.category(c) != 'Mn')


def clean_text(x):
    if not isinstance(x, str):
        return x
    x = html.unescape(x)
    x = re.sub(r'<[^>]+>', ' ', x)          # bỏ thẻ HTML (<p>…)
    x = re.sub(r'\s+', ' ', x).strip()
    return x or None


def parse_date(s):
    if isinstance(s, (pd.Timestamp, datetime)):
        return pd.Timestamp(s)
    if not isinstance(s, str) or not s.strip():
        return pd.NaT
    s = s.strip().replace('-', '/')
    for f in ('%d/%m/%Y %H:%M', '%d/%m/%Y'):
        try:
            return pd.to_datetime(s, format=f)
        except ValueError:
            pass
    return pd.NaT


def vnd(x):
    return f"{x:,.0f}".replace(',', '.')


def to_num(s):
    try:
        return float(s.replace(',', ''))
    except (ValueError, AttributeError):
        return None



def canon_vendor(v):
    return VENDOR_MAP.get(v.strip().upper(), v.strip().title()) if v else None


def is_solution(seg):
    toks = [t for t in re.split(r'[^a-z0-9]+', fold(seg)) if t]
    if not toks:
        return False
    hit = sum(1 for t in toks if t in SOLUTION_KW or t in FILLER or t.isdigit())
    return hit >= max(1, len(toks) * 0.6)


def customer_from_deal_name(name):
    """Trả về (tên khách hàng, cách lấy)"""
    n = name.strip().rstrip(';').strip()
    f = fold(n)
    if TEST_NAMES.search(f):
        return n, 'test'
    if re.match(r'^C12-BCA 175 - ', n):                         # phần sau là reseller
        return 'C12-BCA 175', 'tach'
    mc = re.search(r'\bcho\s+(.+)$', n)
    if PROJECT_PREFIX.match(f) and mc:                          # "Mua sắm … cho X" -> X
        return mc.group(1).strip(), 'tach'
    if '_' in n:                                                # "KhachHang_Hang/GiaiPhap"
        parts = [p.strip() for p in n.split('_') if p.strip()]
        keep = [p for p in parts if not is_solution(p)]
        if keep and len(keep) < len(parts):
            return ' - '.join(keep), 'tach'
        if len(keep) == len(parts):
            return ' - '.join(parts), 'giu'
        return n, 'kiem tra'
    m = re.match(r'^(.*\S)\s*-\s*([^-]+)$', n)                  # "KhachHang - Hang"
    if m and is_solution(m.group(2)):
        return m.group(1).strip(), 'tach'
    if PROJECT_PREFIX.match(f):
        return n, 'kiem tra'
    return n, 'giu'


def deal_size(v):
    if v is None or pd.isna(v) or v <= 0:
        return None
    for lim, lab in DEAL_SIZE:
        if v < lim:
            return lab
    return '> 5B'



# Nhận diện dòng BOM là dịch vụ khi cột Hãng trống (so trên Tên giải pháp + Ghi chú đã bỏ dấu)
SERVICE_RX = r'^\s*dich vu|trien khai|cai dat|dao tao'
OTHER_RX = r'giai phap khac|san pham khac|hang khac'
NON_MAIN = ('Khác', 'Dịch vụ')
SRC_RANK = {'BOM': 0, 'Tên giải pháp/Ghi chú': 1, 'Tên hãng (CRM)': 2, 'Từ khóa': 3, 'Không xác định': 4}


def num(v):
    """Số từ ô BOM dạng text ('10,500', '119,746,000'). Trả về (giá trị, lỗi?)"""
    if v is None or (not isinstance(v, str) and pd.isna(v)):
        return None, False
    s = str(v).strip()
    if s.startswith('#'):                      # #VALUE!, #N/A… = lỗi công thức CRM
        return None, True
    try:
        return float(s.replace(',', '')), False
    except ValueError:
        return None, True


def keyword_vendors(text):
    """Các hãng tìm thấy theo từ khóa, theo thứ tự xuất hiện"""
    f = fold(text)
    found = sorted((m.start(), v) for v, rx in VENDOR_KEYWORDS for m in [re.search(rx, f)] if m)
    out = []
    for _, v in found:
        if v not in out:
            out.append(v)
    return out


def deal_level_vendors(r):
    """Suy hãng ở cấp deal: Tên hãng (CRM) -> Từ khóa. Trả về ([hãng], nguồn)"""
    col = canon_vendor(r['Tên hãng']) if isinstance(r['Tên hãng'], str) else None
    if col and col != 'Khác':
        return [col], 'Tên hãng (CRM)'
    text = ' '.join(str(r[c]) for c in ['Thông tin BOM (model, số lượng, năm support...)', 'Tên thương vụ', 'Tên dự án']
                    if isinstance(r[c], str))
    kw = keyword_vendors(text)
    if kw:
        return kw, 'Từ khóa'
    if col == 'Khác':
        return ['Khác'], 'Tên hãng (CRM)'
    return [], 'Không xác định'


def line_vendor(line, deal_row):
    """Hãng của 1 dòng BOM. Trả về (hãng | None, nguồn)"""
    h = line[BOM_COLS['Hãng']]
    if isinstance(h, str) and h.strip():
        return canon_vendor(h), 'BOM'
    text = ' '.join(str(line[BOM_COLS[c]]) for c in ['Tên giải pháp', 'Ghi chú'] if isinstance(line[BOM_COLS[c]], str))
    if text:
        kw = keyword_vendors(text)
        if kw:
            return kw[0], 'Tên giải pháp/Ghi chú'
        if re.search(SERVICE_RX, fold(text)):
            return 'Dịch vụ', 'Tên giải pháp/Ghi chú'
        if re.search(OTHER_RX, fold(text)):
            return 'Khác', 'Tên giải pháp/Ghi chú'
    vs, src = deal_level_vendors(deal_row)
    return (vs[0], src) if vs else (None, 'Không xác định')


def read_source(src):
    xl = pd.ExcelFile(src)
    sheets = [s for s in xl.sheet_names if s.startswith(SHEET_PREFIX)]
    if not sheets:
        sys.exit(f'Không tìm thấy sheet bắt đầu bằng "{SHEET_PREFIX}". Các sheet có: {xl.sheet_names}')
    return xl.parse(sheets[0]), sheets[0]


def bom_lines(raw):
    """Tách sheet nguồn thành (bảng deal, bảng dòng BOM)"""
    raw = raw.copy()
    raw['_row'] = raw.index + 2                                   # số dòng Excel để truy vết
    raw['_ID'] = raw['ID'].ffill()
    deals = raw[raw['ID'].notna()].copy()
    cont = raw[raw['ID'].isna()]
    deal_cols = [c for c in raw.columns if c not in BOM_COLS.values() and not c.startswith('_')]
    leaked = cont[deal_cols].notna().any(axis=1)
    if leaked.any():
        print(f'CẢNH BÁO: {leaked.sum()} dòng BOM nối tiếp (ID trống) có dữ liệu ở cột thông tin deal – '
              f'kiểm tra dòng Excel {cont.loc[leaked, "_row"].tolist()[:10]}')
    if raw['_ID'].isna().any():
        sys.exit('Sheet nguồn có dòng BOM nằm trước dòng deal đầu tiên (không xác định được ID) – dừng để kiểm tra.')
    lines = raw[raw[BOM_CORE].notna().any(axis=1)].copy()
    return deals, lines


# ============================== DANH MỤC SALE ==============================
# Tài khoản CRM -> tên đầy đủ. Sale mới: thêm 1 dòng vào đây.
SALE_NAMES = {
    'thuydinh': 'Đinh Văn Thủy', 'hangnguyen': 'Nguyễn Thu Hằng', 'trinhluong': 'Lương Thị Tuyết Trinh',
    'phamquyet': 'Phạm Văn Quyết', 'theduong': 'Dương Văn Thế', 'quyvy': 'Vy Công Quý',
    'kienkhong': 'Khổng Đức Kiên', 'nhungnguyen': 'Nguyễn Hồng Nhung', 'hungle': 'Lê Huy Hùng',
    'anhle': 'Lê Tuấn Anh', 'thangnguyen': 'Nguyễn Duy Thắng', 'nghiemphan': 'Phan Quế Nghiệm',
    'hungpham': 'Phạm Quốc Hùng', 'tupham': 'Phạm Cao Anh Tú', 'hieudinh': 'Đinh Văn Hiệu',
    'hiendoan': 'Đoàn Thị Hiền', 'hungtran': 'Trần Quang Hưng',
    'truongdang': 'Đặng Đình Trường',
}


def _sale_tags():
    """Các kiểu sale hay ghi kèm vào tên KH: _KienKhong, _QuyetPham, _Tuấn Anh, _Lê Hùng, _Trinh…"""
    tags = set()
    for user, full in SALE_NAMES.items():
        p = fold(full).split()
        fam, mid, giv = p[0], p[1:-1], p[-1]
        tags |= {user, fam + giv, giv + fam, giv, ''.join(p), ''.join(mid[-1:]) + giv}
        if len(p) >= 2:
            tags.add(''.join(p[-2:]))
    return {t for t in tags if len(t) >= 4}


SALE_TAGS = None


def strip_sale_tag(name):
    m = re.match(r'^(.*\S)\s*_\s*([^_]+)$', name)
    global SALE_TAGS
    SALE_TAGS = SALE_TAGS or _sale_tags()
    if m and re.sub(r'[^a-z]', '', fold(m.group(2))) in SALE_TAGS:
        return m.group(1).strip(), m.group(2).strip()
    return name, None


# ============================== PHÂN NHÓM KHÁCH HÀNG ==============================
SEGMENTS = ['Cơ quan Đảng, Nhà nước', 'Bộ, Ngành, Cơ quan TW', 'UBND / Sở ban ngành địa phương',
            'An ninh - Quốc phòng', 'Ngân hàng - Tài chính - Bảo hiểm', 'Viễn thông - CNTT',
            'Năng lượng - Điện - Dầu khí', 'Y tế - Giáo dục', 'Sản xuất - Thương mại - Dịch vụ',
            'Giao thông - Vận tải - Logistics', 'DNNN khác', 'Khác']
# Quy tắc áp dụng theo THỨ TỰ, khớp quy tắc đầu tiên thì dừng (so trên tên KH đã bỏ dấu, chữ thường).
SEGMENT_RULES = [
    ('Viễn thông - CNTT', r'viettel|\bvtg\b|natcom'),                       # Viettel xét trước (xem ngoại lệ bên dưới)
    ('An ninh - Quốc phòng', r'cong an|\bbca\b|\bbo ca\b|^ca |\bca (tp|tinh|thanh pho)\b|\bpa0\d|'
                             r'\b(a0[1-9]|c0[1-9]|c1[0-2]|c500|h0[1-9]|v0[1-9]|b01|x01)\b|quoc phong|\bbqp\b|quan doi|'
                             r'quan khu|bo tu lenh|binh chung|bien phong|canh sat|quan y|hai quan|military|'
                             r'tong cuc (2|ii|tinh bao|ky thuat|chinh tri)|co yeu|si quan|tong tham muu|tac chien|'
                             r'an ninh nhan dan|bo tham muu|bo doi'),
    ('Cơ quan Đảng, Nhà nước', r'tinh uy|thanh uy|quan uy|huyen uy|\bvptu\b|trung uong dang|tw dang|vptw|'
                               r'tap chi cong san|phu chu tich|van phong chu tich nuoc|quoc hoi|toa an|vien kiem sat|'
                               r'kiem toan nha nuoc'),
    ('UBND / Sở ban ngành địa phương', r'\bubnd\b|uy ban nhan dan|^so |\bso (khoa|kh|thong tin|tttt|4t|xay dung|'
                                       r'dan toc|tai chinh|y te|giao duc|noi vu|tu phap|ke hoach|cong thuong|van hoa)\b|'
                                       r'bqlda|ban quan ly khu kinh te|ban duy tu|chuyen doi so|'
                                       r'phuc vu hanh chinh cong|dai phat thanh|bao va phat thanh'),
    ('Bộ, Ngành, Cơ quan TW', r'^bo |\bbo (noi vu|tai chinh|tai nguyen|ngoai giao|dan toc|xay dung|khoa hoc|'
                              r'cong thuong|y te|giao duc|tu phap|ke hoach|nong nghiep|giao thong|van hoa|thong tin)\b|'
                              r'\bcuc\b|tong cuc|kho bac|bao hiem xa hoi|\bbhxh\b|ngan hang nha nuoc|'
                              r'(uy ban|ub) chung khoan|uy ban tieu chuan|quy phat trien|thong tin tin dung|'
                              r'\bcic\b|vnnic|ttxvn|thong tan xa|vien han lam'),
    ('Năng lượng - Điện - Dầu khí', r'\bevn|dien luc|thuy dien|nhiet dien|truyen tai dien|he thong dien|'
                                    r'thi truong dien|dieu do|\bnpt\b|\bptc ?\d|^pc |genco|\bpvn\b|pv ?gas|petro|'
                                    r'dau khi|khi viet nam|^than |\btkv\b|vimico|khoang san|nhien lieu|xang dau|nsrp|'
                                    r'npcit|\bnpc\b|\bspc\b'),
    ('Y tế - Giáo dục', r'benh vien|^bv|\bbvdk\b|y te|huyet hoc|^truong |dai hoc|hoc vien|giao duc|educa|'
                        r'\bttyt\b|hospital'),
    ('Ngân hàng - Tài chính - Bảo hiểm', r'ngan hang|bank|chung khoan|securities|bao hiem|insurance|tai chinh vi mo|'
                                         r'finance|napas|vnpay|thanh toan|\b(bidv|vietcombank|vietinbank|agribank|'
                                         r'mbbank|tpbank|vib|scb|shb|shbfc|lpb|lpbs|exim|eximbank|tcbs|hsc|kbsv|'
                                         r'vndirect|mbs|vps|shs|dnse|dsc|agriseco|opes|pjico|pvi|insmart|mfinance|'
                                         r'timo|oceanbank|publicbank)\b'),
    ('Giao thông - Vận tải - Logistics', r'hang khong|airline|airport|san bay|\bacv\b|vaeco|viags|sasco|\bvna\b|'
                                         r'truc thang|\bvnh\b|quan ly bay|vnaic|hang hai|giao hang|\bghtk\b|'
                                         r'logistics|van tai|duong sat|\bvetc\b|\bitl\b|vnpost|buu chinh'),
    ('Viễn thông - CNTT', r'vnpt|mobifone|\bcmc\b|telecom|vien thong|\bfpt\b|\bfis\b|\bctin\b|cong nghe|technology|'
                          r'tecapro|gosu|vtvcab|software|vietbay|tntech'),
    ('DNNN khác', r'cap nuoc|sawaco|vietlott|xo so|trac dia|dong tau'),
    ('Sản xuất - Thương mại - Dịch vụ', r'cong ty|\bcty\b|\bctcp\b|tnhh|tap doan|group|corporation|\bcorp\b|'
                                        r'limited|\bjsc\b|holdings'),
]
# Ngoại lệ đã xác định cụ thể (tên KH bỏ dấu, chữ thường -> nhóm). Thêm KH mới khó phân loại vào đây.
SEGMENT_OVERRIDES = {
    'tong cong ty cp buu chinh viettel': 'Giao thông - Vận tải - Logistics',
    'viettel post': 'Giao thông - Vận tải - Logistics',
    'cong ty thuong mai & xuat nhap khau/ thuong mai & xnk viettel (viettel commerce)': 'Sản xuất - Thương mại - Dịch vụ',
    'cong ty lien doanh thap ngan hang dau tu va phat trien viet nam': 'Sản xuất - Thương mại - Dịch vụ',
    'vietinbank gold & jewellery': 'Ngân hàng - Tài chính - Bảo hiểm',
    'pvi': 'Ngân hàng - Tài chính - Bảo hiểm',
    'cong ty tnhh mot thanh vien dong tau hong ha': 'DNNN khác',
    'cong ty tnhh mtv trac dia ban do': 'DNNN khác',
    'tap doan phenikaa': 'Sản xuất - Thương mại - Dịch vụ',
    'access trade': 'Sản xuất - Thương mại - Dịch vụ', 'career viet': 'Sản xuất - Thương mại - Dịch vụ',
    'diana unicharm': 'Sản xuất - Thương mại - Dịch vụ', 'everland van don': 'Sản xuất - Thương mại - Dịch vụ',
    'fujikin': 'Sản xuất - Thương mại - Dịch vụ', 'freshmart (c.p)': 'Sản xuất - Thương mại - Dịch vụ',
    'fujimart': 'Sản xuất - Thương mại - Dịch vụ', 'jarllytec': 'Sản xuất - Thương mại - Dịch vụ',
    'maison': 'Sản xuất - Thương mại - Dịch vụ', 'masterise': 'Sản xuất - Thương mại - Dịch vụ',
    'messer viet nam': 'Sản xuất - Thương mại - Dịch vụ', 'nutifood': 'Sản xuất - Thương mại - Dịch vụ',
    'panasonic r&d': 'Sản xuất - Thương mại - Dịch vụ', 'rang dong': 'Sản xuất - Thương mại - Dịch vụ',
    'shemar power vn': 'Sản xuất - Thương mại - Dịch vụ', 'sun group': 'Sản xuất - Thương mại - Dịch vụ',
    'yakult viet nam': 'Sản xuất - Thương mại - Dịch vụ', 'yokohama tyre viet nam': 'Sản xuất - Thương mại - Dịch vụ',
    'berjaya gia thinh': 'Sản xuất - Thương mại - Dịch vụ',
    'viet duc co so 3': 'Y tế - Giáo dục',
    'cong ty tnhh mtv van hanh htd va ttd quoc gia': 'Năng lượng - Điện - Dầu khí',
    'gtb endpoint protector': 'Khác', 'giai phap attt cap do 3': 'Khác',
}
# KH chưa đủ thông tin để phân nhóm chắc chắn -> để "Khác" và gắn cờ cần kiểm tra
SEGMENT_UNSURE = {'hal vn', 'mcst', 'itdb', 'lpex', 'plc', 'intech',
                  'trung tam nghien cuu va ung dung cong nghe truyen thong (r&d)'}


def classify_customer(name):
    """Trả về (nhóm KH, chắc chắn?)"""
    f = re.sub(r'\s+', ' ', fold(name)).strip()
    if f in SEGMENT_UNSURE:
        return ('Viễn thông - CNTT' if 'cong nghe' in f else 'Khác'), False
    if f in SEGMENT_OVERRIDES:
        return SEGMENT_OVERRIDES[f], True
    for seg, rx in SEGMENT_RULES:
        if re.search(rx, f):
            if seg == 'Viễn thông - CNTT' and re.search(r'viettel', f):
                if re.search(r'buu chinh|post', f):
                    return 'Giao thông - Vận tải - Logistics', True
                if re.search(r'commerce|thuong mai', f):
                    return 'Sản xuất - Thương mại - Dịch vụ', True
            return seg, True
    return 'Khác', False



# ============================== XỬ LÝ CHÍNH ==============================
def deal_bom(lines_of_deal, deal_row):
    """Đọc các dòng BOM của 1 deal -> danh sách dòng {vendor, src, amount, err, computed, sales, link}"""
    items = []
    for _, ln in lines_of_deal.iterrows():
        amt, err = num(ln[BOM_COLS['Thành tiền']])
        computed = False
        if amt is None:
            q, _ = num(ln[BOM_COLS['Số lượng']])
            p, _ = num(ln[BOM_COLS['Đơn giá']])
            if q is not None and p is not None:
                amt, computed = q * p, True
        v, src = line_vendor(ln, deal_row)
        link = ln[BOM_COLS['Link deal']]
        m = re.search(r'/(\d+)\s*$', link) if isinstance(link, str) else None
        su = ln[BOM_COLS['Sales user']]
        items.append({'vendor': v, 'src': src, 'amount': amt, 'err': err, 'computed': computed,
                      'sales': su.strip().lower() if isinstance(su, str) and su.strip() else None,
                      'link': int(m.group(1)) if m else None})
    return items


def vendor_rows(items, deal_row):
    """Trả về (danh sách dòng output [(hãng, nguồn, giá trị BOM theo hãng)], hãng chính, cách chọn)"""
    if not items:
        vs, src = deal_level_vendors(deal_row)
        return ([(v, src, None) for v in vs] or [(None, 'Không xác định', None)]), (vs[0] if vs else None), None
    order, amt, cnt, srcs = [], {}, {}, {}
    for it in items:
        v = it['vendor']
        if v not in order:
            order.append(v)
        cnt[v] = cnt.get(v, 0) + 1
        if it['amount'] is not None:
            amt[v] = amt.get(v, 0) + it['amount']
        srcs.setdefault(v, []).append(it['src'])
    known = [v for v in order if v is not None]
    pool = [v for v in known if v not in NON_MAIN] or known
    if not pool:
        return [(None, 'Không xác định', amt.get(None))], None, None
    main = max(pool, key=lambda v: (amt.get(v, 0), cnt[v]))
    how = 'giá trị BOM lớn nhất' if amt.get(main, 0) > 0 else 'nhiều dòng BOM nhất (BOM chưa có giá)'
    rest = sorted([v for v in known if v != main], key=lambda v: -amt.get(v, 0))
    if None in order:
        rest.append(None)
    rows = [(v, min(srcs[v], key=SRC_RANK.get), amt.get(v)) for v in [main] + rest]
    return rows, main, how


def is_lead(r, items):
    acts = sum(float(r[c]) for c in ACTIVITY if pd.notna(r[c]))
    v0 = r['Giá trị thương vụ']
    return (r['Trạng thái thương vụ'] == 'Active' and r['Giai đoạn'] == LEAD_STAGE and not items
            and (pd.isna(v0) or v0 == 0) and pd.isna(r['Timeline dự án']) and acts == 0)


def convert(src, dst, ref_date):
    raw, sheet = read_source(src)
    missing = [c for c in REQUIRED if c not in raw.columns]
    if missing:
        sys.exit(f'Sheet "{sheet}" thiếu cột: {missing}. Kiểm tra lại cấu hình export CRM.')
    for c in raw.columns:
        if raw[c].dtype == object or str(raw[c].dtype) in ('str', 'string'):
            raw[c] = raw[c].map(clean_text)
    deals, lines = bom_lines(raw)
    if deals['ID'].duplicated().any():
        sys.exit(f'Sheet nguồn có ID deal trùng: {deals.loc[deals["ID"].duplicated(), "ID"].tolist()[:10]} – dừng để kiểm tra.')
    deals['ID'] = deals['ID'].astype('int64')
    lines['_ID'] = lines['_ID'].astype('int64')
    n_raw_rows, n_deals_raw, n_lines_raw = len(raw), len(deals), len(lines)

    df = deals.copy()
    for c in ['Timeline dự án', 'Ngày ký hợp đồng', 'Từ lúc', 'Cập nhật lần cuối']:
        parsed = df[c].map(parse_date)
        bad = df[c].notna() & parsed.isna()
        if bad.any():
            print(f'CẢNH BÁO: {bad.sum()} giá trị ngày không đọc được ở cột "{c}": {df.loc[bad, c].tolist()[:5]}')
        df[c] = parsed

    # ---- Bước 1: Tên khách hàng + loại deal test (giữ nguyên quy tắc cũ, thêm mẫu "ko sale")
    cust, cust_note, test_rows = [], [], []
    for _, r in df.iterrows():
        if TEST_NAMES.search(fold(r['Tên thương vụ'])):
            test_rows.append(r['ID'])
        if isinstance(r['Tên công ty'], str) and r['Tên công ty']:
            cust.append(r['Tên công ty']); cust_note.append(None)
            continue
        c, how = customer_from_deal_name(r['Tên thương vụ'])
        if how == 'giu' and is_solution(c):
            how = 'kiem tra'
        cust.append(c)
        if how in ('test', 'giu'):
            cust_note.append(None if how == 'test' else 'Tên KH lấy từ tên thương vụ (CRM trống "Tên công ty")')
        elif how == 'tach':
            cust_note.append('Tên KH lấy từ tên thương vụ (CRM trống "Tên công ty"), đã bỏ phần hãng/giải pháp/reseller')
        else:
            cust_note.append('Tên KH lấy nguyên tên thương vụ nhưng tên giống tên dự án/giải pháp – cần kiểm tra lại')
    stripped = [strip_sale_tag(c) for c in cust]
    df['_cust'] = [c for c, _ in stripped]
    df['_sale_tag'] = [t for _, t in stripped]
    df['_cust_note'] = cust_note
    test_names = df.loc[df['ID'].isin(test_rows), 'Tên thương vụ'].tolist()
    df = df[~df['ID'].isin(test_rows)].copy()
    lines = lines[~lines['_ID'].isin(test_rows)]

    df['_key'] = df['_cust'].map(lambda s: re.sub(r'[^a-z0-9]+', ' ', fold(s)).strip())
    def pick(g):
        vc = g.value_counts()
        return sorted(vc.index, key=lambda s: (fold(s) == s.lower(), s.isupper(), -vc[s], s))[0]
    df['_cust_std'] = df['_key'].map(df.groupby('_key')['_cust'].agg(pick))

    by_deal = {i: g for i, g in lines.groupby('_ID')}
    name_counts = df.groupby('Tên thương vụ')['ID'].apply(list)
    out, split_flags = [], []
    st_cnt = {'lines': 0, 'value_err': 0, 'value_fixed': 0, 'unresolved': 0, 'inferred': 0}

    for _, r in df.iterrows():
        items = deal_bom(by_deal.get(r['ID'], lines.iloc[0:0]), r)
        lead = is_lead(r, items)
        rows, main, how = vendor_rows(items, r)
        priced = [it['amount'] for it in items if it['amount'] is not None]
        bom = sum(priced)
        st_cnt['lines'] += len(items)
        issues, notes, lead_keep = [], [], []

        # ---- giá trị (giữ nguyên quy tắc cũ; tổng BOM nay lấy từ cột Thành tiền)
        v0, cur = r['Giá trị thương vụ'], r['Tiền tệ thương vụ']
        if (pd.isna(v0) or v0 == 0) and bom > 0:
            val = bom; notes.append(f'Giá trị lấy từ tổng BOM ({vnd(bom)}) vì giá trị CRM = 0')
        elif pd.isna(v0) or v0 == 0:
            val = None; issues.append('Thiếu giá trị thương vụ (CRM = 0, BOM không có giá)')
        elif cur == 'USD' and v0 < USD_MISLABEL_THRESHOLD:
            val = v0 * USD_RATE; notes.append(f'Quy đổi {v0:,.0f} USD x {USD_RATE:,} (tỷ giá giả định)')
        else:
            val = v0
            if cur == 'USD':
                notes.append('CRM ghi tiền tệ USD nhưng số tiền là VND – giữ số gốc, coi là VND')
        if val and bom > 0 and abs(val - bom) > BOM_TOLERANCE * bom:
            issues.append(f'Giá trị {vnd(val)} lệch tổng BOM {vnd(bom)} ({(val - bom) / bom:+.0%})')

        # ---- chất lượng dòng BOM
        n_err = sum(1 for it in items if it['err'])
        n_fix = sum(1 for it in items if it['computed'])
        n_unpriced = sum(1 for it in items if it['amount'] is None)
        st_cnt['value_err'] += n_err; st_cnt['value_fixed'] += n_fix
        if n_err:
            issues.append(f'BOM có {n_err} dòng Thành tiền lỗi công thức CRM (#VALUE!)')
        if n_fix:
            notes.append(f'{n_fix} dòng BOM thiếu Thành tiền – đã tính = Số lượng x Đơn giá')
        if n_unpriced:
            issues.append(f'BOM có {n_unpriced} dòng chưa có giá')
        blank = [it for it in items if it['src'] != 'BOM']
        if blank:
            st_cnt['inferred'] += len(blank)
            issues.append(f'BOM có {len(blank)} dòng trống cột Hãng')
            res = [it for it in blank if it['vendor']]
            unres = len(blank) - len(res)
            st_cnt['unresolved'] += unres
            if res:
                desc = ', '.join(sorted({f'{it["vendor"]} (suy từ {it["src"]})' for it in res}))
                notes.append(f'Hãng của {len(res)} dòng BOM trống đã suy: {desc}')
            if unres:
                issues.append(f'{unres} dòng BOM không xác định được hãng')
        if not items and not lead:
            issues.append('Deal không có dòng BOM')
            if main:
                notes.append(f'Hãng suy từ {rows[0][1]}')

        # ---- sale, link deal
        u = r['Chủ sở hữu thương vụ']
        uk = u.strip().lower() if isinstance(u, str) else None
        if uk in SALE_NAMES:
            sale = SALE_NAMES[uk]
        else:
            sale = u if isinstance(u, str) else None
            if isinstance(u, str):
                lead_keep.append(f'Sale "{u}" chưa có trong bảng tên sale – giữ tên tài khoản CRM')
        su = sorted({it['sales'] for it in items if it['sales'] and it['sales'] != uk})
        if su:
            issues.append(f'Sales user ở dòng BOM ({", ".join(SALE_NAMES.get(s, s) for s in su)}) khác chủ deal ({sale})')
        links = sorted({it['link'] for it in items if it['link'] and it['link'] != r['ID']})
        if links:
            issues.append(f'Link deal ở dòng BOM trỏ sang deal khác (ID {", ".join(map(str, links))}) – BOM có thể copy từ deal khác')

        # ---- trạng thái / giai đoạn / close date (giữ nguyên)
        st, gd = r['Trạng thái thương vụ'], r['Giai đoạn']
        if st == 'Active' and gd == 'Thực hiện hợp đồng':
            issues.append('Trạng thái Active nhưng giai đoạn CRM là "Thực hiện hợp đồng" (có thể đã Won)')
        if st == 'Won' and gd != 'Thực hiện hợp đồng':
            issues.append(f'Trạng thái Won nhưng giai đoạn CRM là "{gd}"')
        if st == 'Won' and pd.isna(r['Ngày ký hợp đồng']):
            issues.append('Won nhưng thiếu ngày ký hợp đồng')
        if st == 'Lost' and not isinstance(r['Lý do thất bại'], str):
            issues.append('Lost nhưng thiếu lý do thất bại')
        cd = r['Timeline dự án']
        overdue = (ref_date - cd).days if pd.notna(cd) and st == 'Active' and cd < ref_date else None
        if overdue is not None:
            issues.append(f'Close date {cd:%d/%m/%Y} đã qua {overdue} ngày nhưng deal vẫn Active')
        same = [i for i in name_counts[r['Tên thương vụ']] if i != r['ID']]
        if same:
            ids = ', '.join(str(i) for i in same[:5]) + (' …' if len(same) > 5 else '')
            issues.append(f'Trùng tên thương vụ với {len(same)} deal khác (ID {ids})')
        if not isinstance(u, str):
            issues.append('Thiếu Sale (owner) trong CRM')
        miss = [lab for lab, ok in [('Close date', pd.notna(cd)), ('Reseller', isinstance(r['Reseller'], str)),
                                    ('Phân loại HĐ', isinstance(r['Phân loại HĐ'], str))] if not ok]
        if miss:
            issues.append('Thiếu thông tin: ' + ', '.join(miss))
        if main is None and not lead:
            issues.append('Chưa xác định được hãng (BOM, Tên giải pháp/Ghi chú, Tên hãng và từ khóa đều không có)')
        if how and len(rows) > 1:
            notes.append(f'Hãng chính = {main} (theo {how})')

        # ---- tên KH, deal size, phân nhóm, failed (giữ nguyên)
        if isinstance(r['_cust_note'], str):
            (lead_keep if 'cần kiểm tra' in r['_cust_note'] else notes).append(r['_cust_note'])
        if isinstance(r['_sale_tag'], str):
            notes.append(f'Đã bỏ hậu tố tên sale "_{r["_sale_tag"]}" trong tên KH')
        if r['_cust_std'] != r['_cust']:
            notes.append(f'Tên KH chuẩn hóa từ "{r["_cust"]}"')
        ds = deal_size(val)
        ds0 = r['Phân loại nhóm deal']
        if ds and not isinstance(ds0, str):
            notes.append('Deal size tính từ giá trị (CRM trống)')
        elif ds and ds0 != ds:
            notes.append(f'Deal size CRM "{ds0}" không khớp giá trị – đã tính lại')
        seg, sure = classify_customer(r['_cust_std'])
        if not sure:
            lead_keep.append(f'Phân nhóm KH "{seg}" chưa chắc chắn – cần kiểm tra')
        reasons = [x for x in [r['Lý do thất bại'], r['Failed additional comment']] if isinstance(x, str)]
        failed = ' – '.join(reasons) if reasons else None

        if lead:   # deal Lead: chỉ 1 cờ chính + các cờ cần người kiểm tra (sale, phân nhóm, tên KH)
            issues = ['Lead chiến dịch – chưa có thông tin cơ hội (chưa có BOM, giá trị, Close date, hoạt động)'] + lead_keep
        else:
            issues += lead_keep

        n = len(rows)
        for k, (ven, vsrc, vamt) in enumerate(rows):
            extra = []
            if n > 1:
                extra = ([f'Deal có {n} dòng hãng: {", ".join(v or "(trống)" for v, _, _ in rows)} – giá trị thương vụ ghi ở dòng này']
                         if k == 0 else [f'Dòng tách hãng {k + 1}/{n} của deal ID {r["ID"]}; giá trị thương vụ chỉ ghi ở dòng hãng chính ({rows[0][0]})'])
            text = []
            if issues:
                text.append('[VẤN ĐỀ] ' + '; '.join(issues))
            if notes or extra:
                text.append('[ĐÃ XỬ LÝ] ' + '; '.join(extra + notes))
            out.append({
                'ID': r['ID'], 'Tên thương vụ': r['Tên thương vụ'], 'Loại deal': 'Lead chiến dịch' if lead else 'Cơ hội',
                'Tên khách hàng': r['_cust_std'], 'Phân nhóm khách hàng': seg, 'Sale': sale, 'Reseller': r['Reseller'],
                'Giai đoạn': gd, 'Trạng thái': st, 'Giá trị thương vụ': val if k == 0 else None,
                'Hãng': ven, 'Nguồn hãng': vsrc, 'Giá trị BOM theo hãng': vamt,
                'Thông tin BOM (model, số lượng, năm support...)': r['Thông tin BOM (model, số lượng, năm support...)'],
                'Phân loại HĐ': r['Phân loại HĐ'], 'Phân loại deal size': ds, 'Close date': cd,
                'Số ngày quá hạn forecast': overdue, 'Lý do Failed': failed, 'Ngày tạo': r['Từ lúc'],
                'Cập nhật lần cuối': r['Cập nhật lần cuối'], 'Cờ chất lượng': ' | '.join(text) or None,
            })
            split_flags.append(k > 0)
    res = pd.DataFrame(out)

    # ---- Đối soát
    first = res.drop_duplicates('ID')
    stats = {
        'sheet': sheet, 'raw_rows': n_raw_rows, 'deals_raw': n_deals_raw, 'lines_raw': n_lines_raw,
        'tests': len(test_rows), 'test_names': test_names, 'deals': res['ID'].nunique(), 'rows': len(res),
        'lead': int((first['Loại deal'] == 'Lead chiến dịch').sum()), 'co_hoi': int((first['Loại deal'] == 'Cơ hội').sum()),
        'lines_used': st_cnt['lines'], 'lines_after_test': len(lines),
        'value_total': first['Giá trị thương vụ'].sum(), 'value_cohoi': first.loc[first['Loại deal'] == 'Cơ hội', 'Giá trị thương vụ'].sum(),
        'bom_total': res['Giá trị BOM theo hãng'].sum(), **st_cnt,
    }
    assert stats['lines_used'] == stats['lines_after_test'], 'Số dòng BOM đã đọc khác số dòng BOM trong file nguồn'
    assert stats['deals'] == n_deals_raw - len(test_rows), 'Số deal output khác số deal raw trừ deal test'
    write_excel(res, split_flags, dst, ref_date, stats)
    return res, stats


def write_excel(res, split_flags, dst, ref_date, s):
    F = 'Arial'
    wb = Workbook(); ws = wb.active; ws.title = 'Data'
    thin = Side(style='thin', color='D9D9D9'); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    widths = {'ID': 10, 'Tên thương vụ': 42, 'Loại deal': 15, 'Tên khách hàng': 38, 'Phân nhóm khách hàng': 26, 'Sale': 20,
              'Reseller': 18, 'Giai đoạn': 20, 'Trạng thái': 11, 'Giá trị thương vụ': 18, 'Hãng': 15, 'Nguồn hãng': 20,
              'Giá trị BOM theo hãng': 18, 'Thông tin BOM (model, số lượng, năm support...)': 50, 'Phân loại HĐ': 13,
              'Phân loại deal size': 12, 'Close date': 12, 'Số ngày quá hạn forecast': 12, 'Lý do Failed': 40,
              'Ngày tạo': 16, 'Cập nhật lần cuối': 16, 'Cờ chất lượng': 90}
    new_cols = {'Loại deal', 'Nguồn hãng', 'Giá trị BOM theo hãng', 'Số ngày quá hạn forecast'}
    cols = list(res.columns)
    for j, h in enumerate(cols, 1):
        c = ws.cell(1, j, h)
        c.font = Font(name=F, bold=True, color='FFFFFF', size=10)
        c.fill = PatternFill('solid', fgColor='2E75B6' if h in new_cols else '1F4E78'); c.border = bd
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[L(j)].width = widths.get(h, 15)
    ws.row_dimensions[1].height = 32
    split_fill = PatternFill('solid', fgColor='F2F2F2')
    fmt = {'Giá trị thương vụ': '#,##0', 'Giá trị BOM theo hãng': '#,##0', 'Close date': 'dd/mm/yyyy',
           'Ngày tạo': 'dd/mm/yyyy hh:mm', 'Cập nhật lần cuối': 'dd/mm/yyyy hh:mm', 'ID': '0', 'Số ngày quá hạn forecast': '0'}
    for i, (rec, is_split) in enumerate(zip(res.itertuples(index=False), split_flags), 2):
        for j, v in enumerate(rec, 1):
            if v is not None and not isinstance(v, str) and pd.isna(v):
                v = None
            if isinstance(v, pd.Timestamp):
                v = v.to_pydatetime()
            c = ws.cell(i, j, v); c.font = Font(name=F, size=10); c.border = bd
            if cols[j - 1] in fmt:
                c.number_format = fmt[cols[j - 1]]
            if is_split:
                c.fill = split_fill
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:{L(len(cols))}{len(res) + 1}'

    g = wb.create_sheet('Huong_dan')
    lines = [
        ('HƯỚNG DẪN ĐỌC FILE', None),
        ('Ngày chốt số liệu', ref_date.strftime('%d/%m/%Y')),
        ('Nguồn dữ liệu', f'Sheet "{s["sheet"]}" của file export CRM. BOM chỉ lấy từ 9 cột "Bảng danh mục sản phẩm dự án - …"; '
                          'không dùng cột text tổng hợp "Bảng danh mục sản phẩm dự án".'),
        ('Đối soát', f'Raw: {s["raw_rows"]} dòng = {s["deals_raw"]} deal, {s["lines_raw"]} dòng BOM. Loại {s["tests"]} deal test. '
                     f'Output: {s["deals"]} deal ({s["co_hoi"]} Cơ hội, {s["lead"]} Lead chiến dịch), {s["rows"]} dòng; '
                     f'đã đọc {s["lines_used"]} dòng BOM.'),
        ('Quy tắc đếm', 'Một deal nhiều hãng được tách thành nhiều dòng cùng ID. Đếm số deal phải dùng ID duy nhất. '
                        'Giá trị thương vụ chỉ ghi ở dòng hãng chính (dòng tách tô nền xám, giá trị để trống), nên cộng '
                        'cột Giá trị thương vụ không bị đếm trùng.'),
        ('Loại deal', 'Lead chiến dịch = deal Active, giai đoạn "Đăng kí cơ hội", không có dòng BOM, giá trị = 0, không có Close date '
                      'và không có hoạt động (email, họp, gọi, ghi chú, ghi nhanh = 0). Còn lại là Cơ hội. Mặc định phân tích chỉ tính '
                      'Cơ hội. Deal Lead tự chuyển sang Cơ hội khi sale bổ sung một trong các thông tin trên hoặc chuyển giai đoạn.'),
        ('Giá trị thương vụ', 'Đơn vị VND. Lấy giá trị CRM; nếu CRM = 0 thì lấy tổng Thành tiền BOM. Deal gắn USD có giá trị '
                              f'>= {USD_MISLABEL_THRESHOLD:,} coi là VND nhập nhầm nhãn. Để trống = chưa có dữ liệu (không phải 0).'),
        ('Hãng', 'Lấy từ cột "Bảng danh mục sản phẩm dự án - Hãng". Nếu trống thì suy lần lượt từ: Tên giải pháp/Ghi chú của dòng '
                 'đó (dòng dịch vụ gán "Dịch vụ") → cột "Tên hãng" → từ khóa trong Thông tin BOM / Tên thương vụ / Tên dự án. '
                 'Dòng đầu của mỗi deal là hãng chính (hãng có giá trị BOM lớn nhất; "Khác", "Dịch vụ" chỉ là hãng chính khi '
                 'deal không có hãng nào khác). Để trống = chưa xác định được.'),
        ('Nguồn hãng', 'BOM = lấy trực tiếp từ cột Hãng của BOM · Tên giải pháp/Ghi chú = suy từ dòng BOM · Tên hãng (CRM) = cột '
                       '"Tên hãng" của deal · Từ khóa = tìm trong Thông tin BOM / tên thương vụ / tên dự án · Không xác định.'),
        ('Giá trị BOM theo hãng', 'Tổng Thành tiền các dòng BOM của hãng đó (dòng thiếu Thành tiền nhưng có Số lượng và Đơn giá '
                                  'thì tính = SL x ĐG; dòng lỗi #VALUE! coi là chưa có giá). Dùng để xem pipeline theo hãng.'),
        ('Phân loại deal size', '< 1B: dưới 1 tỷ · 1B - 2B: 1 đến dưới 2 tỷ · 2B - 5B: 2 đến dưới 5 tỷ · > 5B: từ 5 tỷ. '
                                'Tính lại từ Giá trị thương vụ, không dùng phân loại CRM vì CRM ghi không nhất quán.'),
        ('Sale', 'Đổi tài khoản CRM (cột "Chủ sở hữu thương vụ") sang họ tên theo bảng ở sheet Danh_muc. '
                 'Tài khoản chưa có trong bảng thì giữ nguyên và gắn cờ. Sales user ở dòng BOM khác chủ deal được gắn cờ.'),
        ('Phân nhóm khách hàng', 'Phân tự động theo tên KH với 12 khối ở sheet Danh_muc. Cơ quan nhà nước phân theo cấp '
                                 '(Đảng/NN, Bộ ngành TW, địa phương); mọi đơn vị thuộc Bộ Công an, Bộ Quốc phòng vào An ninh - '
                                 'Quốc phòng; doanh nghiệp phân theo ngành kinh doanh chính, công ty con theo ngành của tập đoàn. '
                                 'KH chưa phân được hoặc chưa chắc chắn có cờ cần kiểm tra.'),
        ('Giai đoạn / Trạng thái', 'Giữ nguyên giá trị CRM ("Giai đoạn", "Trạng thái thương vụ"); mâu thuẫn giữa hai cột được '
                                   'ghi ở Cờ chất lượng.'),
        ('Close date', 'Lấy từ cột "Timeline dự án" của CRM = ngày dự kiến Won (forecast).'),
        ('Số ngày quá hạn forecast', 'Ngày chốt số liệu − Close date; chỉ tính cho deal Active có Close date đã qua.'),
        ('Lý do Failed', 'Gộp "Lý do thất bại" và "Failed additional comment": <lý do> – <ghi chú>.'),
        ('Cờ chất lượng', '[VẤN ĐỀ] = dữ liệu thiếu/mâu thuẫn cần người phụ trách kiểm tra (luôn ghi lỗi gốc kể cả khi đã suy '
                          'thông tin). [ĐÃ XỬ LÝ] = thay đổi so với CRM (tên KH, hãng, giá trị, deal size). Deal Lead chỉ có cờ '
                          '"Lead chiến dịch" và các cờ cần kiểm tra tên KH / sale / phân nhóm.'),
        ('Tham số', f'Ngưỡng USD nhập nhầm: {USD_MISLABEL_THRESHOLD:,}; tỷ giá giả định: {USD_RATE:,}; '
                    f'ngưỡng lệch BOM: {BOM_TOLERANCE:.0%}.'),
    ]
    for i, (a, b) in enumerate(lines, 1):
        ca = g.cell(i, 1, a); ca.font = Font(name=F, size=10, bold=True); ca.alignment = Alignment(vertical='top')
        cb = g.cell(i, 2, b); cb.font = Font(name=F, size=10); cb.alignment = Alignment(wrap_text=True, vertical='top')
    g['A1'].font = Font(name=F, size=12, bold=True)
    g.column_dimensions['A'].width = 26; g.column_dimensions['B'].width = 110

    dm = wb.create_sheet('Danh_muc')
    hdr = PatternFill('solid', fgColor='1F4E78')
    for j, h in enumerate(['Tài khoản CRM', 'Tên Sale', '', 'Phân nhóm khách hàng', 'Số deal Cơ hội', 'Ví dụ khách hàng'], 1):
        if h:
            c = dm.cell(1, j, h); c.font = Font(name=F, bold=True, color='FFFFFF', size=10); c.fill = hdr
    for i, (u, n) in enumerate(SALE_NAMES.items(), 2):
        dm.cell(i, 1, u).font = Font(name=F, size=10); dm.cell(i, 2, n).font = Font(name=F, size=10)
    first = res.drop_duplicates('ID')
    first = first[first['Loại deal'] == 'Cơ hội']
    for i, seg in enumerate(SEGMENTS, 2):
        sub = first[first['Phân nhóm khách hàng'] == seg]
        ex = ', '.join(sub['Tên khách hàng'].value_counts().index[:4])
        for j, v in ((4, seg), (5, len(sub)), (6, ex)):
            dm.cell(i, j, v).font = Font(name=F, size=10)
    for col, w in zip('ABCDEF', (16, 24, 3, 32, 12, 90)):
        dm.column_dimensions[col].width = w
    wb.save(dst)


def ref_from_name(path):
    m = re.search(r'_(\d{2})_(\d{2})_(\d{2})\.xlsx$', path)
    if m:
        try:
            return pd.Timestamp(datetime.strptime('/'.join(m.groups()), '%d/%m/%y'))
        except ValueError:
            pass
    return pd.Timestamp.today().normalize()



if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('src'); ap.add_argument('dst'); ap.add_argument('--ref-date')
    a = ap.parse_args()
    ref = pd.Timestamp(datetime.strptime(a.ref_date, '%d/%m/%Y')) if a.ref_date else ref_from_name(a.src)
    res, s = convert(a.src, a.dst, ref)
    print(f'Ngày chốt số liệu: {ref:%d/%m/%Y} | Sheet nguồn: {s["sheet"]}')
    print(f'Raw: {s["raw_rows"]} dòng = {s["deals_raw"]} deal, {s["lines_raw"]} dòng BOM')
    print(f'Loại {s["tests"]} deal test: {s["test_names"]}')
    print(f'Output: {s["deals"]} deal ({s["co_hoi"]} Cơ hội, {s["lead"]} Lead chiến dịch), {s["rows"]} dòng')
    print(f'Dòng BOM đã đọc: {s["lines_used"]} | trống Hãng: {s["inferred"]} (không suy được: {s["unresolved"]}) | '
          f'lỗi #VALUE!: {s["value_err"]} | tính lại SL x ĐG: {s["value_fixed"]}')
    print(f'Tổng giá trị thương vụ: {s["value_total"]:,.0f} VND (Cơ hội: {s["value_cohoi"]:,.0f}) | '
          f'Tổng giá trị BOM theo hãng: {s["bom_total"]:,.0f} VND')
